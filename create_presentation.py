import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)  # 16:9 widescreen
    prs.slide_height = Inches(7.5)

    blank_slide_layout = prs.slide_layouts[6]
    
    # Palette
    NAVY = RGBColor(26, 54, 93)       # #1A365D Primary dark
    BLUE = RGBColor(43, 108, 176)     # #2B6CB0 Accent blue
    DARK_GRAY = RGBColor(45, 55, 72)  # #2D3748 Body text
    LIGHT_GRAY = RGBColor(113, 128, 150) # #718096 Secondary text
    BG_LIGHT = RGBColor(247, 250, 252) # #F7FAFC Card background
    WHITE = RGBColor(255, 255, 255)
    GREEN = RGBColor(40, 167, 69)     # #28A745 Success
    RED = RGBColor(220, 53, 69)       # #DC3545 Malignant
    BORDER_COLOR = RGBColor(226, 232, 240)

    base_dir = r"C:\Users\rajee\Desktop\ml-intern-assignment"
    results_dir = os.path.join(base_dir, "results")
    
    def add_header(slide, title_text, category_text="SUPERVISED MACHINE LEARNING PIPELINE"):
        # Header banner text
        tx_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.1))
        tf = tx_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = BLUE
        p_cat.space_after = Pt(2)
        
        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = NAVY

    def add_card(slide, left, top, width, height, bg_color=BG_LIGHT, border_color=BORDER_COLOR):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1)
        return shape

    # ==========================================
    # SLIDE 1: Title Slide
    # ==========================================
    slide1 = prs.slides.add_slide(blank_slide_layout)
    
    # Background card
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = NAVY
    bg1.line.fill.background()

    # Inner decorative container
    inner = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(0.8), Inches(11.733), Inches(5.9))
    inner.fill.solid()
    inner.fill.fore_color.rgb = RGBColor(33, 66, 110)
    inner.line.color.rgb = RGBColor(66, 110, 160)
    inner.line.width = Pt(1.5)

    tx1 = slide1.shapes.add_textbox(Inches(1.4), Inches(1.3), Inches(10.5), Inches(4.8))
    tf1 = tx1.text_frame
    tf1.word_wrap = True
    
    p = tf1.paragraphs[0]
    p.text = "COMPUTER SCIENCE / AI INTERN ASSIGNMENT"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = RGBColor(99, 179, 237)
    p.space_after = Pt(14)
    
    p = tf1.add_paragraph()
    p.text = "Breast Cancer Diagnostic Classification"
    p.font.size = Pt(32)
    p.font.bold = True
    p.font.color.rgb = WHITE
    p.space_after = Pt(4)

    p = tf1.add_paragraph()
    p.text = "An End-to-End, Interpretable Supervised Machine Learning Pipeline"
    p.font.size = Pt(18)
    p.font.color.rgb = RGBColor(226, 232, 240)
    p.space_after = Pt(28)

    p = tf1.add_paragraph()
    p.text = "Candidate: Intern Candidate   |   Domain: Clinical Oncology & AI Diagnostics   |   Stack: Python, Scikit-Learn"
    p.font.size = Pt(11)
    p.font.color.rgb = RGBColor(203, 213, 225)
    p.space_after = Pt(16)
    
    p = tf1.add_paragraph()
    p.text = "Key Result: 98.25% Test Accuracy  •  98.61% F1-Score  •  97.62% Malignant Sensitivity  •  Strict Zero Data Leakage"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = RGBColor(104, 211, 145)

    # ==========================================
    # SLIDE 2: Problem Statement & Clinical Relevance
    # ==========================================
    slide2 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide2, "1. Problem Statement & Clinical Relevance")
    
    # Left Box: Problem Context
    add_card(slide2, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tx2_l = slide2.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tf2_l = tx2_l.text_frame
    tf2_l.word_wrap = True
    
    p = tf2_l.paragraphs[0]
    p.text = "The Clinical Challenge"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)
    
    bullets_l = [
        ("High Mortality Risk:", " Early and accurate detection of breast cancer is critical for patient survival and timely surgical/oncological intervention."),
        ("Biopsy Digitization:", " When a breast lesion is detected, clinicians perform a Fine Needle Aspirate (FNA). High-resolution digital microscopy extracts nuclear morphological measurements."),
        ("Diagnostic Objective:", " Automate the classification of cell nuclei into Malignant (cancerous) vs. Benign (non-cancerous) to assist pathologists with rapid, reliable triage.")
    ]
    for b_title, b_desc in bullets_l:
        p = tf2_l.add_paragraph()
        p.text = "• " + b_title
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        p_run = p.add_run()
        p_run.text = b_desc
        p_run.font.bold = False
        p.space_after = Pt(10)

    # Right Box: ML Problem Formulation
    add_card(slide2, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tx2_r = slide2.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tf2_r = tx2_r.text_frame
    tf2_r.word_wrap = True
    
    p = tf2_r.paragraphs[0]
    p.text = "Machine Learning Formulation"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)
    
    bullets_r = [
        ("Task Category:", " Supervised Binary Classification (two mutually exclusive diagnostic classes)."),
        ("Target Variable:", " 'target' encoded as 0 = Malignant (cancerous), 1 = Benign (safe)."),
        ("Input Predictors:", " 30 continuous real-valued cellular morphology measurements."),
        ("Why ML?:", " Reduces human diagnostic fatigue, eliminates inter-observer variance, and provides instant calibrated probabilities (<1 ms inference)."),
        ("Cost Asymmetry:", " False Negatives (missing cancer) are clinically far more dangerous than False Positives.")
    ]
    for b_title, b_desc in bullets_r:
        p = tf2_r.add_paragraph()
        p.text = "• " + b_title
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        p_run = p.add_run()
        p_run.text = b_desc
        p_run.font.bold = False
        p.space_after = Pt(8)

    # ==========================================
    # SLIDE 3: Dataset Selection & Integrity Audit
    # ==========================================
    slide3 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide3, "2. Dataset Selection & Data Integrity Audit")
    
    # 3 Stat Cards on Top
    stat_data = [
        ("569 Biopsies", "Total Samples (UCI Benchmark)", Inches(0.8)),
        ("30 Features", "Continuous Nuclear Traits", Inches(4.8)),
        ("0 Missing / 0 Duplicates", "100% Data Integrity", Inches(8.8))
    ]
    for val, lbl, left_pos in stat_data:
        add_card(slide3, left_pos, Inches(1.6), Inches(3.7), Inches(1.2), bg_color=RGBColor(235, 248, 255), border_color=BLUE)
        tx = slide3.shapes.add_textbox(left_pos + Inches(0.1), Inches(1.7), Inches(3.5), Inches(1.0))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = val
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.alignment = PP_ALIGN.CENTER
        p2 = tf.add_paragraph()
        p2.text = lbl
        p2.font.size = Pt(10)
        p2.font.color.rgb = DARK_GRAY
        p2.alignment = PP_ALIGN.CENTER

    # Bottom Container
    add_card(slide3, Inches(0.8), Inches(3.0), Inches(11.7), Inches(3.8))
    tx3_b = slide3.shapes.add_textbox(Inches(1.1), Inches(3.2), Inches(11.1), Inches(3.4))
    tf3_b = tx3_b.text_frame
    tf3_b.word_wrap = True
    
    p = tf3_b.paragraphs[0]
    p.text = "Dataset Architecture & Feature Decomposition"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(8)
    
    p = tf3_b.add_paragraph()
    p.text = "• Dataset Provenance: Collected by Dr. William H. Wolberg, W. Nick Street, and Olvi L. Mangasarian (University of Wisconsin-Madison). Distributed via UCI ML Repository & Scikit-Learn."
    p.font.size = Pt(11)
    p.font.color.rgb = DARK_GRAY
    p.space_after = Pt(6)

    p = tf3_b.add_paragraph()
    p.text = "• 10 Core Morphological Attributes: Radius, Texture, Perimeter, Area, Smoothness, Compactness, Concavity, Concave Points, Symmetry, and Fractal Dimension."
    p.font.size = Pt(11)
    p.font.color.rgb = DARK_GRAY
    p.space_after = Pt(6)

    p = tf3_b.add_paragraph()
    p.text = "• 3 Statistical Aggregations: For each of the 10 attributes, 3 metrics are recorded: Mean (10), Standard Error (10), and Worst/Largest Value (10) = 30 Input Predictors."
    p.font.size = Pt(11)
    p.font.color.rgb = DARK_GRAY
    p.space_after = Pt(6)

    p = tf3_b.add_paragraph()
    p.text = "• Target Class Balance: 357 Benign (62.74%) vs. 212 Malignant (37.26%). Mild natural imbalance accounted for via Stratified Splitting."
    p.font.size = Pt(11)
    p.font.bold = True
    p.font.color.rgb = BLUE

    # ==========================================
    # SLIDE 4: Exploratory Data Analysis & Visualizations
    # ==========================================
    slide4 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide4, "3. Exploratory Data Analysis & Morphological Visuals")
    
    # Left Image: Distribution Plot
    img_dist = os.path.join(results_dir, "feature_distribution.png")
    if os.path.exists(img_dist):
        slide4.shapes.add_picture(img_dist, Inches(0.8), Inches(1.6), Inches(6.0), Inches(3.6))
        
    # Right Image: Correlation Heatmap
    img_corr = os.path.join(results_dir, "correlation_heatmap.png")
    if os.path.exists(img_corr):
        slide4.shapes.add_picture(img_corr, Inches(7.1), Inches(1.6), Inches(5.4), Inches(3.6))

    # Bottom Takeaways
    add_card(slide4, Inches(0.8), Inches(5.4), Inches(11.7), Inches(1.6))
    tx4_b = slide4.shapes.add_textbox(Inches(1.0), Inches(5.5), Inches(11.3), Inches(1.4))
    tf4_b = tx4_b.text_frame
    tf4_b.word_wrap = True
    
    p = tf4_b.paragraphs[0]
    p.text = "Key EDA Insights:"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(2)
    
    p = tf4_b.add_paragraph()
    p.text = "1. Morphological Divergence: Malignant cell nuclei display significantly higher grayscale texture variance and severe contour concave indentations (Worst Concave Points: 0.182 vs 0.074)."
    p.font.size = Pt(10)
    p.font.color.rgb = DARK_GRAY
    
    p = tf4_b.add_paragraph()
    p.text = "2. Multicollinearity: Nuclear size dimensions (radius, perimeter, area) exhibit near-perfect correlation (r > 0.98), necessitating regularized modeling."
    p.font.size = Pt(10)
    p.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 5: Data Preprocessing & Leakage Prevention
    # ==========================================
    slide5 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide5, "4. Preprocessing & Prevention of Data Leakage")
    
    prep_cards = [
        ("1. Train / Test Partitioning", 
         "• 80% Training Set (455 samples)\n• 20% Unseen Test Set (114 samples)\n• Applied random_state=42 for exact mathematical reproducibility.",
         Inches(0.8)),
        ("2. Stratified Sampling", 
         "• Used stratify=y in train_test_split\n• Preserves exact 62.7% Benign : 37.3% Malignant ratio across train & test splits\n• Eliminates partition bias in evaluation.",
         Inches(4.8)),
        ("3. Leak-Free Standardization", 
         "• Features range from 0.04 (concavity) to 2501 (area)\n• StandardScaler fitted strictly on X_train\n• X_test transformed using training mean & std\n• Zero data leakage guaranteed.",
         Inches(8.8))
    ]
    for title, desc, left_pos in prep_cards:
        add_card(slide5, left_pos, Inches(1.6), Inches(3.7), Inches(5.2))
        tx = slide5.shapes.add_textbox(left_pos + Inches(0.2), Inches(1.8), Inches(3.3), Inches(4.8))
        tf = tx.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = NAVY
        p.space_after = Pt(10)
        
        for line in desc.split("\n"):
            p_line = tf.add_paragraph()
            p_line.text = line
            p_line.font.size = Pt(10.5)
            p_line.font.color.rgb = DARK_GRAY
            p_line.space_after = Pt(6)

    # ==========================================
    # SLIDE 6: Model Selection & Justification
    # ==========================================
    slide6 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide6, "5. Model Selection & Theoretical Justification")
    
    # Left Card: Selected Model
    add_card(slide6, Inches(0.8), Inches(1.6), Inches(5.7), Inches(5.2), bg_color=RGBColor(240, 249, 255), border_color=BLUE)
    tx6_l = slide6.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tf6_l = tx6_l.text_frame
    tf6_l.word_wrap = True
    
    p = tf6_l.paragraphs[0]
    p.text = "Selected Model: Logistic Regression (L2)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)
    
    reasons = [
        ("Calibrated Probabilities:", " Maps linear combinations to posterior confidence scores via sigmoid σ(z), essential for clinical risk triage."),
        ("Intrinsic Interpretability:", " Model weights (w_i) represent direct log-odds, enabling transparent feature influence auditing."),
        ("Convexity & Stability:", " Guarantees a unique global minimum with zero stochastic variance across runs."),
        ("L2 Regularization (C=1.0):", " Ridge penalty constrains weights, preventing overfitting on collinear dimensions.")
    ]
    for r_title, r_desc in reasons:
        p = tf6_l.add_paragraph()
        p.text = "✔ " + r_title
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = DARK_GRAY
        p_run = p.add_run()
        p_run.text = r_desc
        p_run.font.bold = False
        p.space_after = Pt(8)

    # Right Card: Alternative Considered
    add_card(slide6, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tx6_r = slide6.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tf6_r = tx6_r.text_frame
    tf6_r.word_wrap = True
    
    p = tf6_r.paragraphs[0]
    p.text = "Alternative Considered: Decision Tree"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)
    
    alt_points = [
        ("Why Considered:", " Decision trees provide intuitive if-then decision rules and handle non-linear relationships natively."),
        ("Why Rejected for Primary Pipeline:", " Single decision trees suffer from high variance (overfitting) and produce uncalibrated step-function outputs rather than smooth probabilities."),
        ("Comparison Verdict:", " Regularized Logistic Regression achieved superior test generalization (98.25%) and calibrated probabilities required for clinical deployment.")
    ]
    for a_title, a_desc in alt_points:
        p = tf6_r.add_paragraph()
        p.text = "• " + a_title
        p.font.bold = True
        p.font.size = Pt(10.5)
        p.font.color.rgb = DARK_GRAY
        p_run = p.add_run()
        p_run.text = a_desc
        p_run.font.bold = False
        p.space_after = Pt(10)

    # ==========================================
    # SLIDE 7: Experimental Results & Confusion Matrix
    # ==========================================
    slide7 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide7, "6. Quantitative Results & Confusion Matrix")
    
    # Left: Confusion Matrix Image
    img_cm = os.path.join(results_dir, "confusion_matrix.png")
    if os.path.exists(img_cm):
        slide7.shapes.add_picture(img_cm, Inches(0.8), Inches(1.6), Inches(5.2), Inches(4.8))

    # Right: Results Table & Metrics
    add_card(slide7, Inches(6.3), Inches(1.6), Inches(6.2), Inches(5.2))
    tx7_r = slide7.shapes.add_textbox(Inches(6.5), Inches(1.8), Inches(5.8), Inches(4.8))
    tf7_r = tx7_r.text_frame
    tf7_r.word_wrap = True
    
    p = tf7_r.paragraphs[0]
    p.text = "Performance on Unseen Test Set (N=114)"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)
    
    metrics_list = [
        ("Accuracy", "98.25%", "112 / 114 correct diagnoses"),
        ("Precision (Benign)", "98.61%", "71 true benign / 72 predicted benign"),
        ("Recall (Benign)", "98.61%", "71 true benign / 72 actual benign"),
        ("F1-Score", "98.61%", "Harmonic mean of precision & recall"),
        ("Malignant Recall (Sensitivity)", "97.62%", "41 / 42 malignant cancer cases caught")
    ]
    for m_name, m_val, m_note in metrics_list:
        p = tf7_r.add_paragraph()
        p.text = f"{m_name}:  "
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        
        p_val = p.add_run()
        p_val.text = m_val + "  "
        p_val.font.bold = True
        p_val.font.color.rgb = GREEN
        
        p_note = p.add_run()
        p_note.text = f"({m_note})"
        p_note.font.bold = False
        p_note.font.size = Pt(9.5)
        p_note.font.color.rgb = LIGHT_GRAY
        p.space_after = Pt(6)

    p_cm = tf7_r.add_paragraph()
    p_cm.text = "\nConfusion Matrix Breakdown:"
    p_cm.font.bold = True
    p_cm.font.size = Pt(11)
    p_cm.font.color.rgb = NAVY
    p_cm.space_after = Pt(2)
    
    p_cm_sub = tf7_r.add_paragraph()
    p_cm_sub.text = "• TN = 41 (Malignant as Malignant)    • TP = 71 (Benign as Benign)\n• FP = 1 (Malignant as Benign)          • FN = 1 (Benign as Malignant)"
    p_cm_sub.font.size = Pt(9.5)
    p_cm_sub.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 8: Feature Influence & Error Root-Cause
    # ==========================================
    slide8 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide8, "7. Feature Influence & Error Root-Cause Analysis")
    
    # Left Image: Feature Importance
    img_fi = os.path.join(results_dir, "feature_importance.png")
    if os.path.exists(img_fi):
        slide8.shapes.add_picture(img_fi, Inches(0.8), Inches(1.6), Inches(6.0), Inches(5.2))

    # Right Card: Error Analysis
    add_card(slide8, Inches(7.1), Inches(1.6), Inches(5.4), Inches(5.2))
    tx8_r = slide8.shapes.add_textbox(Inches(7.3), Inches(1.8), Inches(5.0), Inches(4.8))
    tf8_r = tx8_r.text_frame
    tf8_r.word_wrap = True
    
    p = tf8_r.paragraphs[0]
    p.text = "Top Predictors & Error Diagnostics"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(8)
    
    p = tf8_r.add_paragraph()
    p.text = "Top Influential Features (Log-Odds):"
    p.font.bold = True
    p.font.size = Pt(11)
    p.font.color.rgb = DARK_GRAY
    
    p_feats = tf8_r.add_paragraph()
    p_feats.text = "1. worst texture (-1.26)\n2. radius error (-1.08)\n3. worst concave points (-0.95)\n4. worst area (-0.95)\n(Negative weights increase log-odds of Malignancy 0)"
    p_feats.font.size = Pt(9.5)
    p_feats.font.color.rgb = BLUE
    p_feats.space_after = Pt(8)

    p_err = tf8_r.add_paragraph()
    p_err.text = "Root-Cause of the 2 Errors (1.75% error rate):"
    p_err.font.bold = True
    p_err.font.size = Pt(11)
    p_err.font.color.rgb = DARK_GRAY
    
    p_err_desc = tf8_r.add_paragraph()
    p_err_desc.text = "• Sample #73 (False Positive / True Malignant): A borderline tumor with atypically smooth cell contours near the 50% boundary (p = 0.537).\n• Sample #541 (False Negative / True Benign): A benign mass exhibiting severe nuclear texture granularity that mimicked malignant dysplasia."
    p_err_desc.font.size = Pt(9.5)
    p_err_desc.font.color.rgb = DARK_GRAY

    # ==========================================
    # SLIDE 9: Conclusion & Engineering Roadmap
    # ==========================================
    slide9 = prs.slides.add_slide(blank_slide_layout)
    add_header(slide9, "8. Conclusion & Production Roadmap")
    
    add_card(slide9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2))
    tx9_l = slide9.shapes.add_textbox(Inches(1.1), Inches(1.8), Inches(5.0), Inches(4.8))
    tf9_l = tx9_l.text_frame
    tf9_l.word_wrap = True
    
    p = tf9_l.paragraphs[0]
    p.text = "Project Summary & Milestones"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)
    
    milestones = [
        ("Robust ML Pipeline:", " Built a fully modular, reproducible supervised classification pipeline in Python."),
        ("Clinical Performance:", " Achieved 98.25% test accuracy and 97.62% malignant sensitivity."),
        ("Full Interpretability:", " Dissected linear log-odds coefficients to validate clinical alignment."),
        ("Serialized Serving:", " Exported trained model and scaler via joblib for sub-millisecond inference.")
    ]
    for m_t, m_d in milestones:
        p = tf9_l.add_paragraph()
        p.text = "✔ " + m_t
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        p_r = p.add_run()
        p_r.text = m_d
        p_r.font.bold = False
        p.space_after = Pt(8)

    add_card(slide9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2))
    tx9_r = slide9.shapes.add_textbox(Inches(7.1), Inches(1.8), Inches(5.1), Inches(4.8))
    tf9_r = tx9_r.text_frame
    tf9_r.word_wrap = True
    
    p = tf9_r.paragraphs[0]
    p.text = "Future Production Improvements"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = NAVY
    p.space_after = Pt(10)
    
    future_work = [
        ("Threshold Tuning:", " Lower classification threshold from 0.50 to 0.30 to prioritize 100% malignant recall."),
        ("Repeated Stratified K-Fold:", " Implement 10-fold cross-validation to assess variance across folds."),
        ("Collinear Feature Elimination:", " Drop redundant perimeter/radius features to simplify model dimensionality."),
        ("Multi-Center Hospital Trials:", " Test generalization on external clinical hospital cohorts.")
    ]
    for f_t, f_d in future_work:
        p = tf9_r.add_paragraph()
        p.text = "➔ " + f_t
        p.font.bold = True
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        p_r = p.add_run()
        p_r.text = f_d
        p_r.font.bold = False
        p.space_after = Pt(8)

    # Save presentation
    ppt_path = os.path.join(base_dir, "Breast_Cancer_ML_Presentation.pptx")
    prs.save(ppt_path)
    print(f"[SUCCESS] Presentation saved successfully at: {ppt_path}")
    return ppt_path

if __name__ == "__main__":
    create_presentation()
